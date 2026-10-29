"""投稿用附表（英文 Excel）：Extended Data Table 1 + Supplementary Tables 1–4。

（分析决策时间线 STable5 只作内部记录，不投稿 —— 2026-10-24 作者定。）

    python tools/export_submission_tables.py --out paper_final/submission

数据一律读 tools/final_tables.py 生成的 paper_final/tables/*.csv（与图、方法与结果对照表同源），
只做翻译与排版：表头加粗冻结、长文本换行、纯数字存成数值、表下附英文注释。
用词与图注一致：overturned / confirmed / too few gold turns；fine-tuned / same-base control；不用 d1–d5 代号。
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
TABLES = REPO / "paper_final" / "tables"
HAN = re.compile(r"[一-鿿]")

# ---------------------------------------------------------------- 词表
OUTCOME = {
    "任一（d1–d4）": "Below minimum standard", "d1 有害内容": "Harmful content",
    "d2 未给可执行安全步骤": "No actionable safety step", "d3 未识别风险": "Risk not recognized",
    "d4 谄媚附和": "Sycophancy", "d5 敷衍回应": "Perfunctory reply",
    "只计 d1/d3/d4": "Dangerous failures (harm, risk or sycophancy)",
}
VERDICT = {
    "一致": "Confirmed", "幅度偏大": "Overturned: judge magnitude too large",
    "幅度偏小": "Overturned: judge magnitude too small", "相反": "Overturned: opposite direction",
    "被否定": "Overturned: not supported", "无法判断": "Too few gold turns to tell",
    "漏掉": "Missed by judge", "都不显著": "Neither significant",
}
YESNO = {"是": "Yes", "否": "No"}
MODEL = {
    "cpsycoun": "CPsyCounX", "internlm2-base": "InternLM2-chat-7B", "mindchat": "MindChat-Qwen2-4B",
    "qwen15-4b-base": "Qwen1.5-4B-Chat", "soulchat": "SoulChat-R1-Qwen3-8B", "qwen3-8b-base": "Qwen3-8B",
}
COMPARISON = {
    "API（含陪伴）− 本地": "Flagship + companion vs local", "心理微调组 − 对照基座组": "Fine-tuned vs control (groups)",
    "旗舰 − 本地": "Flagship vs local", "陪伴 − 旗舰": "Companion vs flagship", "陪伴 − 本地": "Companion vs local",
    "高 − 低风险段（fixed 臂）": "High vs low risk (fixed arm)", "高 − 低风险段（两臂合并）": "High vs low risk (both arms)",
    "高 − 低风险段（本地模型）": "High vs low risk (local models)",
}
GROUP_TOKENS = [  # ST3 子组：长的先替换
    ("组｜本地开源（心理微调 + 对照基座）", "Access: local open models (fine-tuned + control)"),
    ("组｜API（含陪伴）", "Access: API (incl. companion)"),
    ("组×风险段｜API", "Access × risk band: API"), ("组×风险段｜本地", "Access × risk band: local"),
    ("臂×风险段｜", "Arm × risk band: "), ("设计组｜对照基座", "Design group: same-base control"),
    ("设计组｜心理微调", "Design group: fine-tuned"), ("设计组｜旗舰", "Design group: flagship"),
    ("设计组｜陪伴", "Design group: companion"), ("风险段｜", "Risk band: "), ("臂｜", "Arm: "), ("模型｜", "Model: "),
    (" · 高", ", high"), (" · 中", ", mid"), (" · 低", ", low"),
    ("高", "high"), ("中", "mid"), ("低", "low"), ("全部", "All turns"), (" − ", " vs "),
]


def comparison(s: str) -> str:
    if s in COMPARISON:
        return COMPARISON[s]
    if " − " in s and not HAN.search(s):
        a, b = s.split(" − ")
        return f"{MODEL.get(a, a)} vs {MODEL.get(b, b)}"
    return s


def subgroup(s: str) -> str:
    for a, b in GROUP_TOKENS:
        s = s.replace(a, b)
    if " vs " in s:                                   # 差值行：去掉重复的前缀，写成「A vs B」
        x, y = s.split(" vs ", 1)
        px, py = (x.split(": ", 1)[0] if ": " in x else ""), (y.split(": ", 1)[0] if ": " in y else "")
        xs, ys = x.split(": ", 1)[-1], y.split(": ", 1)[-1]
        if px == py:
            s = f"{px}: {xs} vs {ys}"
        else:
            s = f"{xs[0].upper() + xs[1:]} vs {ys}"
    return s


def outcome_phrase(s: str) -> str:
    """「旗舰 − 本地，任一」→「Flagship vs local, below minimum standard」。"""
    if "，" not in s:
        return s
    left, right = s.split("，", 1)
    left = {"旗舰 − 本地": "Flagship vs local", "心理微调 − 对照基座": "Fine-tuned vs control"}.get(left, left).replace(
        " − ", " vs ")
    right = OUTCOME.get(right, {"任一": "below minimum standard"}.get(right, right))
    return f"{left}, {right[0].lower() + right[1:] if right[0].isupper() and right != right.upper() else right}"


# ---------------------------------------------------------------- 各表
def read(name):
    with (TABLES / f"{name}.csv").open(encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def ed1():
    rows = read("EDTable1_target_models")
    head = ["Model", "Model ID", "Group", "Fine-tuned from", "Parameters", "Access", "Decoding", "Pair",
            "Gold-standard turns", "Failed gold turns (below minimum standard)"]
    acc = {"未公开": "Not disclosed", "API（OpenRouter）": "API (OpenRouter)", "本地（vLLM）": "Local (vLLM)",
           "本地（vLLM，关思考）": "Local (vLLM, thinking disabled)",
           "本地（vLLM；分词器指向 CPsyCounX，md5 相同）": "Local (vLLM; tokenizer from CPsyCounX, identical files)"}
    out = []
    for r in rows[1:]:
        model, mid, _, grp, base, size, access, dec, pair, n, nf = r
        pair = re.sub(r"配对 (\d)（(.+?) ↔ (.+?)）", r"Pair \1 (\2 and \3)", pair)
        out.append([model, mid, grp, base, acc.get(size, size), acc.get(access, access), dec.replace("，", ", "), pair, n, nf])
    notes = [
        "Groups were fixed at the design stage. Local open models comprise the three mental-health fine-tuned models "
        "and their same-base controls.",
        "Each pair consists of a mental-health fine-tuned model and the general model it was built from, of the same "
        "size; the two differ mainly in mental-health fine-tuning.",
        "MiniMax-M2-her was called directly, without a character or role-play configuration, like the other models.",
        "Gold-standard turns and failures are unweighted counts from the 180 psychiatrist-rated turns.",
    ]
    return ("Extended Data Table 1 | Target models", head, out, notes)


def st1():
    rows = read("STable1_flip_table")
    head = ["Judge", "Comparison", "Failure type", "Judge estimate (whole corpus)", "Judge 95% CI",
            "Psychiatrist estimate (original weights)", "95% CI (original)",
            "Judge minus psychiatrist (original)", "95% CI of difference (original)", "Holm P of difference (original)",
            "Verdict (original weights)",
            "Psychiatrist estimate (raking)", "95% CI (raking)",
            "Judge minus psychiatrist (raking)", "95% CI of difference (raking)", "Holm P of difference (raking)",
            "Verdict (raking)", "Judge significant", "Overturned under both weightings"]
    out = []
    for r in rows[1:]:
        j, c, o, je, jci, pe, pci, d0, dci0, hp0, v0, pr, prci, d1, dci1, hp1, v1, sig, both = r
        out.append([j, comparison(c), OUTCOME.get(o, o), je, jci, pe, pci, d0, dci0, hp0, VERDICT.get(v0, v0),
                    pr, prci, d1, dci1, hp1, VERDICT.get(v1, v1), YESNO.get(sig, sig), YESNO.get(both, both)])
    notes = [
        "Differences are the first group minus the second, in percentage points. Judge estimates are uncorrected "
        "observations on every turn the judge scored, as a typical evaluation would report them. Psychiatrist "
        "estimates are weighted to the whole corpus. All 95% CIs, including those of the judge-minus-psychiatrist "
        "difference, come from one joint bootstrap that resampled the 16 personas 10,000 times and recomputed the "
        "weights each time.",
        "Verdicts apply to conclusions the judge reached (95% CI excluding zero). Confirmed: psychiatrists significant "
        "in the same direction and the CI of the judge-minus-psychiatrist difference includes zero. Overturned: "
        "psychiatrists significant in the opposite direction, or the CI of the difference excludes zero (not supported "
        "when psychiatrists are not significant; magnitude too large or too small when they are significant in the same "
        "direction). Too few gold turns to tell: psychiatrists not significant and the CI of the difference includes "
        "zero. Missed by judge: judge not significant, psychiatrists significant. Neither significant: neither estimate "
        "excludes zero. Holm P values are adjusted across the cells in which each judge was significant.",
        "Overturned under both weightings may include cells in which the judge itself was not significant but the "
        "psychiatrists were significant in the opposite direction (see the Judge significant column).",
        "The DeepSeek judge scored only the six local open models, so it contributes five comparisons.",
    ]
    return ("Supplementary Table 1 | Judge versus psychiatrist conclusions for all comparisons", head, out, notes)


def st2():
    rows = read("STable2_primary_tests")
    head = ["ID", "Family", "Comparison", "Weights", "Group A turns / failed", "Group B turns / failed",
            "Group A failure rate (%)", "Group B failure rate (%)", "Difference (percentage points)", "95% CI",
            "Bootstrap P", "Holm P within family", "Holm P across families",
            "Significant after within-family Holm", "Degenerate bootstrap"]
    fam = {"PA 模型比较": "PA: model comparison", "PB 心理健康微调": "PB: mental-health fine-tuning"}
    wts = {"原权重": "Original", "raking": "Raking", "截断 raking（稳健性，不参与结论）": "Trimmed raking (robustness only)"}
    out = []
    for r in rows[1:]:
        r = list(r)
        r[1] = fam.get(r[1], r[1])
        c = r[2]
        if c.startswith(("CPsyCounX", "MindChat", "SoulChat")):
            a, rest = c.split(" − ")
            b, oc = rest.split("，")
            c = f"{a} vs {b}, below minimum standard" if oc == "任一" else c
        else:
            c = outcome_phrase(c)
        r[2] = c
        r[3] = wts.get(r[3], r[3])
        r[16] = YESNO.get(r[16], r[16])
        r[17] = YESNO.get(r[17], r[17])
        del r[14:16]                                  # 旧方法 Holm（2 列）
        del r[11]                                     # 旧方法 P
        out.append(r)
    notes = [
        "The hypotheses and test families were not preregistered. PA, flagship versus local open models for falling below the minimum standard. PB, mental-health "
        "fine-tuning: fine-tuned models pooled against their controls (below minimum standard, harmful content, "
        "sycophancy) and each of the three pairs (below minimum standard).",
        "Bootstrap P inverts the percentile bootstrap, 2 × min(P(d* ≤ 0), P(d* ≥ 0)) with a (1 + count)/(B + 1) "
        "correction; the smallest attainable value is about 0.001.",
        "Holm correction was applied within PA and within PB, and across PA and PB together as a stricter check. "
        "Conclusions require significance under both original and raking weights; trimmed raking is a robustness check "
        "only. Resampling: 16 personas, 2,000 bootstrap samples, seed 0.",
    ]
    return ("Supplementary Table 2 | Primary hypothesis tests on psychiatrist ratings", head, out, notes)


def st3():
    rows = read("STable3_gold_prevalence")
    head = ["Type", "Subgroup or comparison", "Failure type", "Gold-standard turns", "Failed turns",
            "Original weights", "95% CI (original)", "Raking", "95% CI (raking)", "Trimmed raking", "95% CI (trimmed)"]
    typ = {"比例": "Failure rate (%)", "差值（百分点）": "Difference (percentage points)"}
    out = []
    for r in rows[1:]:
        r = list(r)
        r[0] = typ.get(r[0], r[0])
        r[1] = subgroup(r[1])
        r[2] = OUTCOME.get(r[2], r[2])
        out.append(r)
    notes = [
        "Psychiatrist consensus ratings of 180 gold-standard turns, weighted to the whole corpus of 9,491 scored turns. "
        "95% CIs from a bootstrap that resampled the 16 personas 2,000 times, recomputing the weights in each sample.",
        "A reply falls below the minimum standard if it contains harmful content, offers no actionable safety step, "
        "misses the risk or is sycophantic. Perfunctory reply is counted separately. Dangerous failures exclude the "
        "safety-step criterion.",
        "Original weights are the ratio of corpus to gold-standard turns within each sampling cell. Raking also calibrates "
        "the weights to the known corpus distribution by sampling cell, model and arm × risk band; about 30% of "
        "bootstrap samples fall back to coarser stratum margins. Original-weight estimates do not use raking. Trimmed raking caps weights at five times the median and is a "
        "robustness check only.",
        "Access groups split models by how they were run (API, including the companion model, versus local); design "
        "groups follow the study design.",
    ]
    return ("Supplementary Table 3 | Psychiatrist-rated failure rates and differences under three weighting schemes",
            head, out, notes)


def st4():
    rows = read("STable4_judge_error_by_model")
    head = ["Judge", "Model", "Failure type", "Gold-standard turns", "TP", "FN (missed)", "FP (false alarm)", "TN",
            "False positive rate", "False negative rate", "Psychiatrist weighted failure rate (%)",
            "Judge weighted failure rate (%)", "O (failures flagged by judge)", "E (expected)", "O/E", "O/E 95% CI",
            "Exact P", "Holm P (60 cells per judge)", "Significant after Holm"]
    out = []
    for r in rows[1:]:
        r = list(r)
        r[2] = OUTCOME.get(r[2], r[2])
        r[18] = YESNO.get(r[18], r[18])
        out.append(r)
    notes = [
        "Counts are unweighted. TP, psychiatrists and judge both rate a failure; FN, psychiatrists rate a failure that "
        "the judge misses; FP, the judge flags a failure the psychiatrists do not. False positive rate = FP/(FP + TN); "
        "false negative rate = FN/(TP + FN).",
        "O is the number of turns the judge rated as failures for that model; E is the number expected from a logistic "
        "error model with the psychiatrist rating, log2 reply length and API access as predictors, fitted to all 180 "
        "turns with sampling weights. O/E above 1 means more failures flagged than expected (false alarms); below 1, "
        "fewer (missed failures). When O = 0 the bootstrap CI is degenerate and should not be interpreted.",
        "O/E 95% CIs from 1,000 bootstrap samples. Exact P from a Poisson-binomial test; Holm correction across the "
        "60 cells (10 models × 6 failure types) of each judge.",
    ]
    return ("Supplementary Table 4 | Judge errors by model", head, out, notes)


# ---------------------------------------------------------------- 写 Excel
NUM = re.compile(r"^[+\-−]?\d+(\.\d+)?(e[+\-]?\d+)?$", re.I)
THIN = Side(style="thin", color="BFBFBF")
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
WRAP_MIN = 28


def to_cell(v: str):
    """纯数字存成数值，并记下原来的小数位（写表时按此设数字格式，-4.0 不会显示成 -4）。"""
    v = v.strip()
    if NUM.match(v):
        x = float(v.replace("−", "-"))
        if re.match(r"^[+\-−]?\d+$", v):
            return int(x), "0"
        if "e" in v.lower():
            return x, "0.0E+00"
        d = len(v.split(".")[1])
        return x, "0." + "0" * d
    return v, None


def write_sheet(ws, title, head, rows, notes):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=THIN, top=THIN)
    for r in rows:
        cells = [to_cell(x) for x in r]
        ws.append([c[0] for c in cells])
        for j, (_, fmt) in enumerate(cells, start=1):
            if fmt:
                ws.cell(ws.max_row, j).number_format = fmt
    widths = []
    for j in range(len(head)):
        vals = [str(head[j])] + [str(r[j]) for r in rows]
        widths.append(min(60, max(10, max(len(v) for v in vals) + 2)))
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
        for row in ws.iter_rows(min_row=4, min_col=j, max_col=j):
            for c in row:
                c.alignment = Alignment(wrap_text=w >= WRAP_MIN, vertical="top")
    ws.freeze_panes = "A4"
    r0 = ws.max_row + 2
    ws.cell(r0, 1, "Notes").font = Font(bold=True)
    for k, n in enumerate(notes, start=1):
        c = ws.cell(r0 + k, 1, n)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r0 + k, start_column=1, end_row=r0 + k, end_column=min(len(head), 10))
        ws.row_dimensions[r0 + k].height = 15 * max(1, len(n) // 140 + 1)


def check_no_chinese(tables):
    bad = []
    for title, head, rows, notes in tables:
        for x in [title, *head, *notes] + [c for r in rows for c in r]:
            if HAN.search(str(x)):
                bad.append((title.split("|")[0].strip(), str(x)[:60]))
    return bad


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="paper_final/submission")
    a = ap.parse_args(argv)
    out = Path(a.out) if Path(a.out).is_absolute() else REPO / a.out
    out.mkdir(parents=True, exist_ok=True)
    ed = ed1()
    sts = [st1(), st2(), st3(), st4()]
    bad = check_no_chinese([ed] + sts)
    if bad:
        for b in bad[:30]:
            print("未翻译：", b)
        return 1
    wb = Workbook()
    write_sheet(wb.active, *ed)
    wb.active.title = "Extended Data Table 1"
    wb.save(out / "Extended_Data_Table_1.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Contents"
    ws.append(["Supplementary Tables"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Sheet", "Title", "Rows"])
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for k, t in enumerate(sts, start=1):
        ws.append([f"ST{k}", t[0].split("|", 1)[1].strip(), len(t[2])])
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 8
    for k, t in enumerate(sts, start=1):
        write_sheet(wb.create_sheet(f"ST{k}"), *t)
    wb.save(out / "Supplementary_Tables.xlsx")
    print(out / "Extended_Data_Table_1.xlsx")
    print(out / "Supplementary_Tables.xlsx")
    for t in [ed] + sts:
        print(f"  {t[0].split('|')[0].strip()}: {len(t[2])} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
