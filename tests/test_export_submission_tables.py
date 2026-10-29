"""投稿附表导出：能从 paper_final/tables 的 CSV 生成两个 Excel，且没有残留中文。"""
import pytest
from openpyxl import load_workbook

from tools import export_submission_tables as X
from tools import final_tables


@pytest.fixture(scope="module", autouse=True)
def _tables(tmp_path_factory):
    """paper_final/tables 不存在时（如公开版），先用 final_tables 生成到临时目录。"""
    if (X.TABLES / "STable1_flip_table.csv").exists():
        yield
        return
    out = tmp_path_factory.mktemp("paper_final")
    final_tables.main(["--out", str(out)])
    old, X.TABLES = X.TABLES, out / "tables"
    yield
    X.TABLES = old


def test_export_runs_and_is_english(tmp_path):
    assert X.main(["--out", str(tmp_path)]) == 0
    wb = load_workbook(tmp_path / "Supplementary_Tables.xlsx")
    assert wb.sheetnames == ["Contents", "ST1", "ST2", "ST3", "ST4"]
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                assert not (isinstance(v, str) and X.HAN.search(v)), v
    assert load_workbook(tmp_path / "Extended_Data_Table_1.xlsx").active.max_row > 10


def test_st1_counts_match_fig3a():
    _, _, rows, _ = X.st1()
    # 列：分类（原权重）= 10，判官显著 = 17（2026-10-26 加入判官 − 医师差值 D 的三列）
    q = [r for r in rows if r[0] == "Qwen" and r[17] == "Yes"]
    over = sum(r[10].startswith("Overturned") for r in q)
    conf = sum(r[10] == "Confirmed" for r in q)
    assert (len(q), over, conf) == (49, 14, 14)


def test_subgroup_difference_labels():
    assert X.subgroup("风险段｜高 − 风险段｜低") == "Risk band: high vs low"
    assert X.subgroup("设计组｜旗舰 − 组｜本地开源（心理微调 + 对照基座）") == \
        "Flagship vs local open models (fine-tuned + control)"
